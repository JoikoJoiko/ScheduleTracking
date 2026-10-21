import hashlib
import io
import json
import zipfile
from datetime import time
from openpyxl import load_workbook
from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone
from .models import *
from .services import audit, require_buildings, publish_version, lock_schedule

FIELDS = ['day', 'slot', 'group', 'subgroup', 'subject', 'teacher', 'building', 'room', 'week', 'type', 'start', 'end']
HEADERS = dict(zip(FIELDS, ['День', 'Пара', 'Группа', 'Подгруппа', 'Дисциплина', 'Преподаватель', 'Корпус', 'Аудитория', 'Неделя', 'Тип', 'Начало', 'Окончание']))
DAYS = ['понедельник', 'вторник', 'среда', 'четверг', 'пятница', 'суббота', 'воскресенье']


def normalize(value):
    return ' '.join(str(value or '').strip().split())


def parse_xlsx(data, profile):
    if len(data) > 10 * 1024 * 1024:
        raise ValidationError('Размер файла превышает 10 МБ')
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 50 * 1024 * 1024 or len(archive.infolist()) > 1000:
                raise ValidationError('Слишком большой распакованный файл')
        book = load_workbook(io.BytesIO(data), data_only=False)
    except (zipfile.BadZipFile, OSError, KeyError, ValueError) as exc:
        raise ValidationError('Некорректный XLSX') from exc
    rows, errors = [], []
    for sheet in book:
        if sheet.max_row > 10000 or sheet.max_column > 200:
            raise ValidationError('Лист превышает лимит 10000 строк / 200 столбцов')
        # Expand merged values before parsing, never guess missing teachers or rooms.
        merged = {}
        for area in sheet.merged_cells.ranges:
            value = sheet.cell(area.min_row, area.min_col).value
            for row in range(area.min_row, area.max_row+1):
                for col in range(area.min_col, area.max_col+1):
                    merged[(row,col)] = value
        def cell(r, c):
            value = merged.get((r,c), sheet.cell(r,c).value)
            if isinstance(value, str) and value.startswith('='):
                raise ValidationError(f'{sheet.title}, {r}:{c}: формулы не поддерживаются, вставьте значения')
            return normalize(value)
        if profile.kind == 'rows':
            header_row = int(profile.mapping.get('header_row', 1))
            names = {cell(header_row, c).casefold():c for c in range(1,sheet.max_column+1)}
            mapping = profile.mapping.get('columns', {})
            columns = {key:names.get(str(mapping.get(key, HEADERS[key])).casefold()) for key in FIELDS}
            missing = [HEADERS[k] for k in ['day','slot','group','subject','teacher','building','room'] if not columns[k]]
            if missing:
                errors.append(f'{sheet.title}: отсутствуют столбцы: {", ".join(missing)}')
                continue
            for r in range(header_row+1, sheet.max_row+1):
                row = {key:cell(r,c) if c else '' for key,c in columns.items()}
                if any(row.values()):
                    row['_source'] = f'{sheet.title}:{r}'
                    rows.append(row)
        elif profile.kind == 'blocks':
            config = profile.mapping
            if not config.get('blocks') or not config.get('day_column') or not config.get('slot_column'):
                raise ValidationError('Настройте blocks, day_column, slot_column профиля')
            for r in range(int(config.get('data_row', 3)), sheet.max_row+1):
                for block in config['blocks']:
                    row = {key:cell(r, int(c)) for key,c in block['columns'].items()}
                    if not any(row.values()):
                        continue
                    row.update(day=cell(r, int(config['day_column'])), slot=cell(r, int(config['slot_column'])),
                        group=cell(int(config.get('group_row',1)),int(block['group_column'])),
                        building=block.get('building',''), _source=f'{sheet.title}:{r}')
                    rows.append(row)
    if len(rows) > 5000:
        raise ValidationError('Не более 5000 занятий за один импорт')
    return rows, errors


def resolve(model, **fields):
    # Python casefold also handles Cyrillic on SQLite, whose LIKE is ASCII-only.
    matches = [obj for obj in model.objects.all() if all(normalize(getattr(obj,k)).casefold() == normalize(v).casefold() for k,v in fields.items())]
    if len(matches) != 1:
        raise ValueError(f'{model._meta.verbose_name}: {fields} — найдено {len(matches)}, исправьте справочник или сопоставление')
    return matches[0]


def validate_rows(rows, user):
    payload, errors, seen = [], [], set()
    for index, row in enumerate(rows, 1):
        try:
            day = normalize(row.get('day')).casefold()
            weekday = DAYS.index(day) if day in DAYS else int(day)-1
            if weekday not in range(7):
                raise ValueError('День должен быть 1–7 или названием дня недели')
            slot = TimeSlot.objects.get(sort_order=int(row.get('slot',0)))
            group = resolve(StudentGroup, name=row.get('group',''))
            building = resolve(Building, name=row.get('building',''))
            room = resolve(Room, building_id=building.pk, number=row.get('room',''))
            teacher = resolve(Teacher, name=row.get('teacher',''))
            subject = resolve(Subject, name=row.get('subject',''))
            subgroup = resolve(Subgroup, group_id=group.pk, name=row['subgroup']) if row.get('subgroup') else None
            require_buildings(user, [building.pk, group.building_id])
            week = normalize(row.get('week','all')).casefold() or 'all'
            week = {'каждая':'all','нечётная':'odd','нечетная':'odd','чётная':'even','четная':'even'}.get(week,week)
            if week not in ['all','odd','even']:
                raise ValueError('Неделя: all / odd / even')
            start = time.fromisoformat(row['start']) if row.get('start') else None
            end = time.fromisoformat(row['end']) if row.get('end') else None
            if bool(start) != bool(end) or (start and start >= end):
                raise ValueError('Некорректное время')
            item = {'day_of_week':weekday,'time_slot_id':slot.pk,'group_id':group.pk,'subgroup_id':subgroup.pk if subgroup else None,
                'room_id':room.pk,'teacher_id':teacher.pk,'subject_id':subject.pk,'week_type':week,
                'lesson_type':row.get('type',''), 'start_time':start.isoformat() if start else None, 'end_time':end.isoformat() if end else None}
            key = json.dumps(item, sort_keys=True)
            if key in seen:
                raise ValueError('Повторяющаяся строка занятия')
            seen.add(key)
            payload.append(item)
        except (ValueError, TypeError, TimeSlot.DoesNotExist) as exc:
            errors.append(f'{row.get("_source", index)}: {exc}')
    if not rows:
        errors.append('Не найдено занятий')
    return payload, errors


def preview_import(upload, profile, period, start, end, user):
    if not upload.name.lower().endswith('.xlsx'):
        raise ValidationError('Поддерживаются только файлы .xlsx')
    if not period.start_date <= start <= end <= period.end_date:
        raise ValidationError('Даты вне учебного периода')
    data = upload.read(10*1024*1024+1)
    rows, errors = parse_xlsx(data, profile)
    payload, row_errors = validate_rows(rows, user)
    batch = ImportBatch.objects.create(filename=upload.name, file_hash=hashlib.sha256(data).hexdigest(), profile=profile,
        period=period,effective_from=start,effective_to=end,payload=payload,errors=errors+row_errors,uploaded_by=user)
    audit(user,'preview_import',batch,new={'rows':len(payload),'errors':len(batch.errors)})
    return batch


@transaction.atomic
def confirm_import(batch, user):
    lock_schedule()
    batch = ImportBatch.objects.select_for_update().get(pk=batch.pk)
    if batch.status == 'confirmed':
        return batch.version
    if batch.status != 'preview' or batch.errors:
        raise ValidationError('Импорт отменён или содержит ошибки')
    require_buildings(user, Room.objects.filter(pk__in=[x['room_id'] for x in batch.payload]).values_list('building_id',flat=True))
    prior = ImportBatch.objects.filter(file_hash=batch.file_hash, period=batch.period,
        effective_from=batch.effective_from,effective_to=batch.effective_to,status='confirmed',version__status='published').first()
    if prior:
        raise ValidationError(f'Этот файл уже опубликован (импорт №{prior.pk})')
    version = ScheduleVersion.objects.create(name=f'Импорт {batch.filename}',period=batch.period,
        effective_from=batch.effective_from,effective_to=batch.effective_to)
    # Identical resource rows represent one joint lesson, not a false room conflict.
    joint = {}
    for raw in batch.payload:
        row = dict(raw)
        group, subgroup = row.pop('group_id'), row.pop('subgroup_id')
        for key in ['start_time','end_time']:
            row[key] = time.fromisoformat(row[key]) if row[key] else None
        key = tuple(row.items())
        lesson = joint.get(key)
        if lesson is None:
            lesson = Lesson.objects.create(version=version, **row)
            joint[key] = lesson
        Participant.objects.create(lesson=lesson, group_id=group, subgroup_id=subgroup)
    publish_version(version,user)
    batch.version = version
    batch.status = 'confirmed'
    batch.completed_at = timezone.now()
    batch.save()
    audit(user,'confirm_import',batch,new={'version':version.pk})
    return version
