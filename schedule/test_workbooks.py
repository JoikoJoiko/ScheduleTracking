import io
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from django.core.exceptions import ValidationError
from django.test import SimpleTestCase
from .college_import import parse_college
from .importing import parse_xlsx, HEADERS, FIELDS
from .models import ImportProfile
from .test_support import college_workbook
from .workbooks import load_schedule_workbook


class WorkbookBoundsTests(SimpleTestCase):
    def saved(self,book):
        output=io.BytesIO();book.save(output)
        return output.getvalue()

    def test_first_course_ignores_formatted_columns_past_actual_table(self):
        book=load_workbook(io.BytesIO(college_workbook(group='ТЕСТ-1-26')))
        sheet=book.active
        sheet.cell(182,255).fill=PatternFill('solid',fgColor='DDDDDD')
        sheet.merge_cells('F4:IU4')
        data=self.saved(book)
        self.assertEqual(load_workbook(io.BytesIO(data)).active.max_column,255)
        rows,metadata,errors=parse_college(data)
        self.assertEqual(errors,[])
        self.assertEqual(metadata['groups'][0]['course'],1)
        self.assertEqual(len(rows),2)

    def test_styled_last_excel_cell_does_not_expand_iteration(self):
        book=load_workbook(io.BytesIO(college_workbook()))
        book.active.cell(1048576,16384).fill=PatternFill('solid',fgColor='DDDDDD')
        cleaned=load_schedule_workbook(self.saved(book))
        self.assertEqual(cleaned.active.max_row,24)
        self.assertEqual(cleaned.active.max_column,9)

    def test_formatted_empty_auxiliary_sheet_is_skipped(self):
        book=load_workbook(io.BytesIO(college_workbook()))
        book.worksheets[1].cell(1048576,255).fill=PatternFill('solid',fgColor='DDDDDD')
        rows,metadata,errors=parse_college(self.saved(book))
        self.assertEqual(errors,[])
        self.assertEqual(len(rows),2)

    def test_parenthesized_group_names_are_preserved(self):
        rows,metadata,errors=parse_college(college_workbook(group='ИСр(к)-9-25'))
        self.assertEqual(errors,[])
        self.assertEqual(metadata['groups'][0]['name'],'ИСр(к)-9-25')
        self.assertEqual(metadata['groups'][0]['course'],2)
        self.assertEqual(rows[0]['group'],'ИСр(к)-9-25')

    def test_real_content_outside_column_limit_is_rejected(self):
        book=load_workbook(io.BytesIO(college_workbook()))
        book.active.cell(4,201,'Данные')
        with self.assertRaisesMessage(ValidationError,'200 столбцов'):
            load_schedule_workbook(self.saved(book))

    def test_real_content_outside_row_limit_is_rejected(self):
        book=load_workbook(io.BytesIO(college_workbook()))
        book.active.cell(10001,1,'Данные')
        with self.assertRaisesMessage(ValidationError,'10000 строк'):
            load_schedule_workbook(self.saved(book))

    def test_standard_profile_also_ignores_trailing_formatting(self):
        book=Workbook();sheet=book.active
        sheet.append([HEADERS[key] for key in FIELDS])
        sheet.append(['Понедельник',1,'ТЕСТ-1-26','','Математика','Тестов А.А.','Колледж','101'])
        sheet.cell(20000,255).fill=PatternFill('solid',fgColor='DDDDDD')
        rows,errors=parse_xlsx(self.saved(book),ImportProfile(kind='rows'))
        self.assertEqual(errors,[])
        self.assertEqual(len(rows),1)
        self.assertEqual(rows[0]['subject'],'Математика')
