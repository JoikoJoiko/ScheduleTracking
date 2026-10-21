"""Bound XLSX input by actual values, not Excel's remembered formatting range."""
import io
import re
import zipfile
from xml.etree import ElementTree as ET
from django.core.exceptions import ValidationError
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter, range_boundaries

MAX_ROWS = 10000
MAX_COLUMNS = 200
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def trim_sheet(xml):
    root = ET.fromstring(xml)
    data = root.find('s:sheetData', NS)
    if data is None:
        return xml
    last_row = last_column = 1
    for row in data:
        for cell in row:
            value = cell.find('s:v', NS)
            inline = cell.find('s:is', NS)
            has_value = (value is not None and value.text is not None) or (
                inline is not None and any((node.text or '').strip() for node in inline.iter()))
            if not has_value and cell.find('s:f', NS) is None:
                continue
            column, number = coordinate_from_string(cell.attrib['r'])
            last_row = max(last_row, number)
            last_column = max(last_column, column_index_from_string(column))
    if last_row > MAX_ROWS or last_column > MAX_COLUMNS:
        raise ValidationError(f'Данные листа выходят за пределы {MAX_ROWS} строк / {MAX_COLUMNS} столбцов '
                              f'(заняты {last_row} строк и {last_column} столбцов).')

    # Trim only the in-memory copy. The user's original workbook remains untouched.
    for row in list(data):
        if int(row.attrib['r']) > last_row:
            data.remove(row)
            continue
        for cell in list(row):
            column, _ = coordinate_from_string(cell.attrib['r'])
            if column_index_from_string(column) > last_column:
                row.remove(cell)
    dimension = root.find('s:dimension', NS)
    if dimension is not None:
        dimension.set('ref', f'A1:{get_column_letter(last_column)}{last_row}')
    merges = root.find('s:mergeCells', NS)
    if merges is not None:
        for area in list(merges):
            left, top, right, bottom = range_boundaries(area.attrib['ref'])
            if left > last_column or top > last_row:
                merges.remove(area)
                continue
            right, bottom = min(right, last_column), min(bottom, last_row)
            if left == right and top == bottom:
                merges.remove(area)
            else:
                area.set('ref', f'{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}')
        merges.set('count', str(len(merges)))
    return ET.tostring(root, encoding='utf-8', xml_declaration=True)


def load_schedule_workbook(data, *, data_only=False):
    if len(data) > 10 * 1024 * 1024:
        raise ValidationError('Размер файла превышает 10 МБ')
    try:
        cleaned = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(cleaned, 'w', zipfile.ZIP_DEFLATED) as target:
            entries = source.infolist()
            if len(entries) > 1000 or sum(entry.file_size for entry in entries) > 50 * 1024 * 1024:
                raise ValidationError('Слишком большой распакованный файл')
            for entry in entries:
                content = source.read(entry)
                if re.fullmatch(r'xl/worksheets/[^/]+\.xml', entry.filename):
                    content = trim_sheet(content)
                target.writestr(entry, content)
        cleaned.seek(0)
        return load_workbook(cleaned, data_only=data_only)
    except (zipfile.BadZipFile, ET.ParseError, OSError, KeyError, ValueError) as exc:
        raise ValidationError('Некорректный XLSX') from exc
