import io
import hashlib
from pathlib import Path
import tempfile
from datetime import date,time,timedelta
from importlib import import_module
from types import SimpleNamespace
from unittest.mock import patch
from PIL import Image
from openpyxl import Workbook
from django.apps import apps
from django.test import TestCase,Client,override_settings
from django.core.files.base import ContentFile
from django.core.management import call_command
from django.core.management.base import CommandError
from . import tests as fixtures,services
from .college_import import parse_college,install_bells
from .importing import populate_version
from .models import (TimeSlot,StudentGroup,GroupYear,Subject,CuratorBell,WeekdayBell,ImageImport,Change,ImportProfile,ImportBatch)
from .timing import curator_bells_for


def hours_workbook():
    book=Workbook();s=book.active
    s['A2']='с 20.10.2026 по 31.12.2026'
    for col,name in [(6,'ЯГ-1-26'),(11,'АГ-1-25')]:
        s.cell(6,col,name);s.merge_cells(start_row=6,start_column=col,end_row=6,end_column=col+3)
    for row,day,slot in [(9,'ПОНЕДЕЛЬНИК',1),(11,None,2),(13,None,3),(16,None,4),
                         (27,'СУББОТА',1),(31,'ПОНЕДЕЛЬНИК',1),(33,None,2),(35,None,3),(38,None,4)]:
        if day:s.cell(row,1,day)
        s.cell(row,2,slot)
        s.cell(row,6,'Математика');s.cell(row+1,6,'Тестов А.А.');s.cell(row,9,'101')
        s.merge_cells(start_row=row,start_column=6,end_row=row,end_column=8)
    for row,col,title in [(8,6,'Разговоры о важном'),(15,11,'Кураторский час'),
                           (30,6,'Разговоры о важом'),(37,11,'разговоры о важном')]:
        s.cell(row,col,title+' Тестов А.А.');s.cell(row,col+3,'102')
        s.merge_cells(start_row=row,start_column=col,end_row=row,end_column=col+2)
    output=io.BytesIO();book.save(output)
    return output.getvalue()


@override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'])
class CuratorHoursTests(TestCase):
    setUp=fixtures.ScheduleTests.setUp
    make_lesson=fixtures.ScheduleTests.make_lesson

    def test_excel_hours_before_monday_and_after_third_pair_in_both_weeks(self):
        rows,meta,errors=parse_college(hours_workbook())
        self.assertEqual(errors,[])
        hours=[r for r in rows if r['slot']==0]
        self.assertEqual(len(hours),4)
        self.assertEqual({(r['day_of_week'],r['week_type'],r['curator_shift']) for r in hours},
                         {(0,w,s) for w in ['odd','even'] for s in [1,2]})
        self.assertEqual({r['subject'] for r in hours},{'Разговоры о важном'})
        self.assertEqual({r['teacher'] for r in hours},{'Тестов А.А.'})
        self.assertEqual({r['group'] for r in hours if r['curator_shift']==1},{'ЯГ-1-26'})

    def test_class_hour_uses_its_shift_and_is_not_a_numbered_pair(self):
        zero=TimeSlot.objects.create(name='Разговоры о важном',sort_order=0)
        subject=Subject.objects.create(name='Кураторский час',is_curricular=False)
        event=self.make_lesson(time_slot=zero,subject=subject,curator_shift=2)
        row=next(r for r in services.schedule_on(self.day) if r['id']==event.pk)
        self.assertEqual((row['slot'],row['start'],row['end']),(0,'14:15','14:45'))
        self.assertEqual(row['subject'],'Разговоры о важном')
        self.assertEqual(row['event_kind'],'class_hour')
        response=self.client.get('/api/subjects',{'group':self.group.pk,'date':str(self.day)}).json()
        self.assertNotIn(subject.pk,[s['id'] for s in response['subjects']])
        event.start_time=time(14,20);event.end_time=time(14,40);event.save()
        self.assertEqual(services.lesson_dict(event,self.day)['start'],'14:20')

    def test_both_curator_bells_are_only_in_monday_api(self):
        TimeSlot.objects.create(name='Кураторский час',sort_order=0,start_time=time(8,20),end_time=time(8,55))
        response=self.client.get('/api/bells',{'date':str(self.day)}).json()
        self.assertEqual([(b['start'],b['end']) for b in response['curator_bells']],
                         [('08:20','08:55'),('14:15','14:45')])
        self.assertEqual(len(response['slots']),2)
        self.assertEqual(curator_bells_for(self.day+timedelta(days=1)),[])
        self.client.force_login(self.admin)
        response=self.client.get('/admin/bells/')
        self.assertTrue(all(slot.sort_order>0 for form,slot in response.context['bell_rows']))

    def test_joint_curator_hour_from_different_course_files_is_not_a_resource_conflict(self):
        zero=TimeSlot.objects.create(name='Кураторский час',sort_order=0)
        subject=Subject.objects.create(name='Разговоры о важном')
        self.make_lesson(time_slot=zero,subject=subject,curator_shift=2)
        other=self.make_lesson(time_slot=zero,subject=subject,curator_shift=2,group=self.group2)
        rows=services.schedule_on(self.day)
        self.assertEqual(services.conflicts_for(self.day,rows),[])
        other.room=self.room2;other.save()
        conflicts=services.conflicts_for(self.day,services.schedule_on(self.day))
        self.assertEqual([c['type'] for c in conflicts],['Пересечение преподавателя'])

    def test_curator_bell_editor_and_validation(self):
        self.client.force_login(self.admin)
        payload={'action':'curator','morning_start':'08:15','morning_end':'08:50',
                 'afternoon_start':'14:10','afternoon_end':'14:40'}
        self.assertEqual(self.client.post('/admin/bells/',payload).status_code,302)
        self.assertEqual(CuratorBell.objects.get(shift=1).start_time,time(8,15))
        payload['morning_end']='08:00'
        self.assertContains(self.client.post('/admin/bells/',payload),'Окончание должно быть позже начала')
        self.assertEqual(CuratorBell.objects.get(shift=1).end_time,time(8,50))
        self.client.force_login(self.dispatcher)
        self.assertEqual(self.client.post('/admin/bells/',payload).status_code,403)

    def test_groups_are_ordered_by_course_in_api_and_model_choices(self):
        GroupYear.objects.filter(group=self.group).update(course=4)
        newest=StudentGroup.objects.create(name='ЯГ-1-26',building=self.building,admission_year=2026,graduation_year=2030)
        second=StudentGroup.objects.create(name='АГ-1-25',building=self.building,admission_year=2025,graduation_year=2029)
        for g,c in [(newest,1),(second,2)]:GroupYear.objects.create(group=g,year=self.year,course=c)
        names=[r['name'] for r in self.client.get('/api/groups').json()['results']]
        self.assertEqual(names,[newest.name,second.name,self.group2.name,self.group.name])
        from .operations import PracticeForm
        self.assertEqual(list(PracticeForm().fields['group'].queryset),[newest,second,self.group2,self.group])

    def test_default_migration_preserves_custom_bells_and_monday_override(self):
        migrate=import_module('schedule.migrations.0009_curator_hours_and_group_order').defaults
        self.slot.start_time=time(8);self.slot.end_time=time(9,30);self.slot.save()
        self.slot2.start_time=time(7);self.slot2.end_time=time(8);self.slot2.save()
        fourth=TimeSlot.objects.create(name='4 пара',sort_order=4,start_time=time(13,20),end_time=time(14,50))
        WeekdayBell.objects.create(day_of_week=0,time_slot=fourth,start_time=time(15),end_time=time(16))
        migrate(apps,SimpleNamespace())
        self.slot.refresh_from_db();self.slot2.refresh_from_db()
        self.assertEqual(self.slot.start_time,time(9))
        self.assertEqual(self.slot2.start_time,time(7))
        self.assertEqual(WeekdayBell.objects.get(time_slot=fourth).start_time,time(15))

    def test_repair_matches_original_hash_is_idempotent_and_preserves_regular_lessons(self):
        TimeSlot.objects.create(name='Разговоры о важном',sort_order=0)
        data=hours_workbook();rows,metadata,_=parse_college(data);metadata['college']=True
        profile=ImportProfile.objects.create(name='Колледж',kind='college')
        batch=ImportBatch.objects.create(filename='course.xlsx',file_hash=hashlib.sha256(data).hexdigest(),
            profile=profile,version=self.version,period=self.period,effective_from=self.version.effective_from,
            effective_to=self.version.effective_to,uploaded_by=self.admin,metadata=metadata,status='confirmed')
        with tempfile.TemporaryDirectory() as directory:
            source=Path(directory)/'course.xlsx';source.write_bytes(data)
            for attempt in range(2):
                call_command('repair_class_hours',str(source),batch=batch.pk,stdout=io.StringIO())
                self.assertEqual(self.version.lessons.filter(time_slot__sort_order=0).count(),4)
                self.assertTrue(self.version.lessons.filter(pk=self.lesson.pk).exists())
                membership=GroupYear.objects.get(group__name='ЯГ-1-26',year=self.year)
                if attempt==0:
                    membership.status='archived';membership.save()
                else:
                    self.assertEqual(membership.status,'archived')
            source.write_bytes(b'changed')
            with self.assertRaisesMessage(CommandError,'отличается'):
                call_command('repair_class_hours',str(source),batch=batch.pk,stdout=io.StringIO())
            self.assertEqual(self.version.lessons.filter(time_slot__sort_order=0).count(),4)
            source.write_bytes(data)
            hour=self.version.lessons.filter(time_slot__sort_order=0).first()
            Change.objects.create(target_date=self.day,original=hour,kind='cancel',created_by=self.admin)
            with self.assertRaisesMessage(CommandError,'есть замены'):
                call_command('repair_class_hours',str(source),batch=batch.pk,stdout=io.StringIO())
            self.assertTrue(self.version.lessons.filter(pk=hour.pk).exists())

    def test_repair_restores_missing_groups_without_replacing_existing_regular_rows(self):
        install_bells()
        self.version.effective_from=date(2026,10,20);self.version.effective_to=date(2026,11,2);self.version.save()
        data=hours_workbook();rows,metadata,_=parse_college(data);metadata['college']=True
        profile=ImportProfile.objects.create(name='Колледж',kind='college')
        batch=ImportBatch.objects.create(filename='course.xlsx',file_hash=hashlib.sha256(data).hexdigest(),
            profile=profile,version=self.version,period=self.period,effective_from=self.version.effective_from,
            effective_to=self.version.effective_to,uploaded_by=self.admin,metadata=metadata,status='confirmed')
        with tempfile.TemporaryDirectory() as directory:
            source=Path(directory)/'course.xlsx';source.write_bytes(data)
            for _ in range(2):
                call_command('repair_class_hours',str(source),batch=batch.pk,restore_missing_groups=True,stdout=io.StringIO())
                restored=self.version.lessons.filter(participants__group__name='ЯГ-1-26',time_slot__sort_order__gt=0)
                self.assertEqual(restored.count(),9)
                self.assertTrue(self.version.lessons.filter(pk=self.lesson.pk).exists())

    def test_repair_reuses_joint_regular_lesson_and_does_not_repeat_participants(self):
        payload={key:getattr(self.lesson,key) for key in ['day_of_week','week_type','time_slot_id','start_time','end_time',
            'subject_id','teacher_id','room_id','lesson_type','curator_shift']}
        payload.update(group_id=self.group2.pk,subgroup_id=None)
        for _ in range(2):populate_version(self.version,[payload],reuse_existing=True)
        self.assertEqual(self.version.lessons.count(),1)
        self.assertEqual(self.lesson.participants.count(),2)


@override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'])
class ImageFragmentTests(TestCase):
    setUp=fixtures.ScheduleTests.setUp
    make_lesson=fixtures.ScheduleTests.make_lesson

    def batch(self):
        image=Image.new('RGB',(100,80),'#ccddff');out=io.BytesIO();image.save(out,format='PNG')
        batch=ImageImport(target_date=self.day,layout='table',uploaded_by=self.admin,
            payload=[{'source_box':[10,10,40,30],'row_box':[0,10,100,30],'inferred':True}])
        batch.image.save('test.png',ContentFile(out.getvalue()),save=False);batch.save()
        return batch

    def test_private_cell_and_row_fragments_and_access_control(self):
        with tempfile.TemporaryDirectory() as folder,override_settings(MEDIA_ROOT=folder):
            batch=self.batch();url=f'/admin/schedule/imageimport/{batch.pk}/fragment/0/'
            self.assertEqual(self.client.get(url).status_code,302)
            self.client.force_login(self.admin)
            for query,size in [({},(120,80)),({'context':'1'},(400,80))]:
                response=self.client.get(url,query)
                self.assertEqual(response.status_code,200)
                self.assertEqual(Image.open(io.BytesIO(response.content)).size,size)
                self.assertIn('private',response['Cache-Control'])
            self.assertEqual(self.client.get(url.replace('/0/','/99/')).status_code,404)
            self.client.force_login(self.dispatcher)
            self.assertEqual(self.client.get(url).status_code,404)
            batch.uploaded_by=self.dispatcher;batch.save()
            self.assertEqual(self.client.get(url).status_code,200)

    def test_invalid_crop_and_reanalysis_csrf_and_publication_state(self):
        with tempfile.TemporaryDirectory() as folder,override_settings(MEDIA_ROOT=folder):
            batch=self.batch();self.client.force_login(self.admin)
            batch.payload[0]['source_box']=[-1,0,200,10];batch.save()
            self.assertEqual(self.client.get(f'/admin/schedule/imageimport/{batch.pk}/fragment/0/').status_code,404)
            url=f'/admin/schedule/imageimport/{batch.pk}/review/'
            guarded=Client(enforce_csrf_checks=True);guarded.force_login(self.admin)
            self.assertEqual(guarded.post(url,{'reanalyze':'1'}).status_code,403)
            with patch('schedule.image_import.analyze',return_value=([{'raw':'new'}],[])) as analyze:
                self.assertEqual(self.client.post(url,{'reanalyze':'1'}).status_code,302)
                self.assertEqual(analyze.call_count,1)
                batch.refresh_from_db();self.assertEqual(batch.payload,[{'raw':'new'}])
                self.assertEqual(batch.status,'review');self.assertFalse(Change.objects.exists())
                batch.status='published';batch.save()
                self.client.post(url,{'reanalyze':'1'})
                self.assertEqual(analyze.call_count,1)

    def test_isolated_retry_recovers_difficult_cell_and_keeps_crop_coordinates(self):
        from .image_import import analyze
        image=Image.new('RGB',(241,41),'white')
        xs=list(range(0,241,30));ys=[0,10,40]
        def text(words,box):
            return self.group.name if box[0]<30 else 'Базы дан' if box[0]<60 else 'По расп'
        with patch('schedule.image_import.table_edges',return_value=(xs,ys)), \
             patch('schedule.image_import.run_ocr',return_value=[]), \
             patch('schedule.image_import.in_box',side_effect=text), \
             patch('schedule.image_import.recognize_boxes',return_value=[self.group.name,'Базы дан']), \
             patch('schedule.image_import.recognize_fragments',return_value=['Базы данных Пигачева А.С. 101']) as retry:
            rows,_=analyze(image,'table',self.day)
        self.assertEqual(len(rows),1);self.assertEqual(retry.call_count,1)
        self.assertEqual(rows[0]['ocr_attempts'],2)
        self.assertEqual(rows[0]['source_box'],[32,12,58,38])
        self.assertEqual(rows[0]['row_box'],[0,10,240,40])
        self.assertEqual(rows[0]['teacher_id'],self.teacher.pk)
        self.assertGreater(rows[0]['confidence'],50)

    def test_isolated_retry_has_a_time_budget(self):
        from .image_import import recognize_fragments
        with patch('schedule.image_import.monotonic',side_effect=[0,0,1,26]), \
             patch('schedule.image_import.run_ocr',return_value=[]) as ocr:
            result=recognize_fragments(Image.new('RGB',(100,100),'white'),[(0,0,10,10),(10,10,20,20)])
        self.assertEqual(len(result),1);self.assertEqual(ocr.call_count,1)
        self.assertLessEqual(ocr.call_args.kwargs['timeout'],20)
