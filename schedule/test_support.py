import io
from openpyxl import Workbook


def college_workbook(group='ТЕСТ-1-23',teacher='Тестов А.А.',room='103п3'):
    """Small synthetic file matching the real college's two-week layout."""
    book=Workbook();sheet=book.active;sheet.title='Расписание'
    sheet['A1']='Расписание занятий';sheet['A2']='с 20.10.2026 по 31.12.2026'
    sheet['F6']=group;sheet.merge_cells('F6:I6');sheet['F7']=1;sheet['H7']=2
    for row in [9,23]:
        sheet.cell(row,1,'ПОНЕДЕЛЬНИК');sheet.cell(row,2,1)
        sheet.cell(row,6,'Программирование');sheet.merge_cells(start_row=row,start_column=6,end_row=row,end_column=8)
        sheet.cell(row,9,room);sheet.cell(row+1,6,teacher);sheet.merge_cells(start_row=row+1,start_column=6,end_row=row+1,end_column=9)
    book.create_sheet('Пустой лист')
    output=io.BytesIO();book.save(output)
    return output.getvalue()
