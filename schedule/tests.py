import io
import json
from datetime import date,time,datetime,timedelta
from openpyxl import Workbook
from django.test import TestCase,Client
from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError,PermissionDenied
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.cache import cache
from django.utils import timezone
from .models import *
from . import services as svc
from .importing import HEADERS,FIELDS,preview_import,confirm_import,parse_xlsx,import_report


class ScheduleTests(TestCase):
    def setUp(self):
        cache.clear()
        WeekdayBell.objects.all().delete()
        TimeSlot.objects.all().delete()
        ImportProfile.objects.all().delete()
        self.admin=get_user_model().objects.create_superuser(username='admin',password='local-test-password')
        self.dispatcher=get_user_model().objects.create_user(username='dispatcher',password='local-test-password',is_staff=True)
        self.year=AcademicYear.objects.create(name='2026/2027',start_date=date(2026,9,1),end_date=date(2027,8,31),status='active')
        self.period=Period.objects.create(name='Семестр',year=self.year,start_date=self.year.start_date,end_date=date(2027,1,31))
        self.building=Building.objects.create(name='Корпус 1')
        self.other_building=Building.objects.create(name='Корпус 2')
        StaffScope.objects.create(user=self.dispatcher).buildings.add(self.building)
        self.type=RoomType.objects.create(name='Лаборатория')
        self.room=Room.objects.create(number='101',building=self.building,room_type=self.type,capacity=30)
        self.room2=Room.objects.create(number='102',building=self.building,room_type=self.type,capacity=20)
        self.remote=Room.objects.create(number='201',building=self.other_building,room_type=self.type)
        self.teacher=Teacher.objects.create(name='Пигачева Анна Сергеевна')
        self.teacher2=Teacher.objects.create(name='Смирнов Пётр Иванович')
        self.subject=Subject.objects.create(name='Базы данных')
        self.specialty=Specialty.objects.create(name='Программирование',code='09.02.07')
        self.group=StudentGroup.objects.create(name='ИС-1-23',specialty=self.specialty,building=self.building,admission_year=2023,graduation_year=2027)
        self.group2=StudentGroup.objects.create(name='ИС-2-24',specialty=self.specialty,building=self.building,admission_year=2024,graduation_year=2028)
        for group in [self.group,self.group2]:
            GroupYear.objects.create(group=group,year=self.year,course=3)
        self.sub1=Subgroup.objects.create(name='1',group=self.group)
        self.sub2=Subgroup.objects.create(name='2',group=self.group)
        self.slot=TimeSlot.objects.create(name='1 пара',sort_order=1,start_time=time(8,30),end_time=time(10))
        self.slot2=TimeSlot.objects.create(name='2 пара',sort_order=2,start_time=time(10,10),end_time=time(11,40))
        self.version=ScheduleVersion.objects.create(name='Основная',period=self.period,effective_from=self.period.start_date,effective_to=self.period.end_date,status='published',published_at=timezone.now())
        self.day=date(2026,10,19)
        self.lesson=self.make_lesson()
        self.profile=ImportProfile.objects.create(name='Стандартный')
    def make_lesson(self,**kwargs):
        group=kwargs.pop('group',self.group);subgroup=kwargs.pop('subgroup',None)
        defaults=dict(version=self.version,day_of_week=0,time_slot=self.slot,subject=self.subject,teacher=self.teacher,room=self.room)
        defaults.update(kwargs)
        lesson=Lesson.objects.create(**defaults)
        Participant.objects.create(lesson=lesson,group=group,subgroup=subgroup)
        return lesson
    def change(self,**kwargs):
        values=dict(target_date=self.day,original=self.lesson,kind='replace',room=self.room2,created_by=self.admin)
        values.update(kwargs)
        return Change.objects.create(**values)
    def workbook(self,rows=None):
        book=Workbook();sheet=book.active
        sheet.append([HEADERS[k] for k in FIELDS])
        for row in rows or [['Понедельник',1,self.group.name,'',self.subject.name,self.teacher.name,self.building.name,self.room.number,'all','','','']]:
            sheet.append(row)
        output=io.BytesIO();book.save(output)
        return SimpleUploadedFile('schedule.xlsx',output.getvalue())
    def preview(self,file=None):
        return preview_import(file or self.workbook(),self.profile,self.period,self.period.start_date,self.period.end_date,self.admin)
    def test_regular_schedule(self):
        rows=svc.schedule_on(self.day)
        self.assertEqual(rows[0]['room_id'],self.room.pk)
        self.assertEqual(svc.schedule_on(self.day+timedelta(days=1)),[])
    def test_draft_does_not_change_public_schedule(self):
        change=self.change()
        self.assertEqual(svc.schedule_on(self.day)[0]['room_id'],self.room.pk)
        self.assertEqual(svc.schedule_on(self.day,preview=change)[0]['room_id'],self.room2.pk)
    def test_published_replacement_preserves_base(self):
        change=self.change();svc.publish_change(change,self.admin)
        row=svc.schedule_on(self.day)[0]
        self.assertEqual(row['room_id'],self.room2.pk)
        self.assertEqual(row['original']['room_id'],self.room.pk)
        self.lesson.refresh_from_db();self.assertEqual(self.lesson.room_id,self.room.pk)
        self.assertTrue(AuditLog.objects.filter(action='publish_change').exists())
    def test_latest_revision_and_withdraw(self):
        a=self.change();svc.publish_change(a,self.admin)
        b=self.change(room=self.remote);svc.publish_change(b,self.admin)
        self.assertEqual(svc.schedule_on(self.day)[0]['room_id'],self.remote.pk)
        svc.withdraw_change(b,self.admin)
        self.assertEqual(svc.schedule_on(self.day)[0]['room_id'],self.room2.pk)
    def test_cancellation_frees_room(self):
        change=self.change(kind='cancel');svc.publish_change(change,self.admin)
        row=svc.availability(self.day,time(8,30),time(10),[self.room])[0]
        self.assertEqual(row['status'],'free')
        self.assertEqual(svc.schedule_on(self.day)[0]['status'],'cancelled')
    def test_move_is_visible_at_destination_only_once(self):
        destination=self.day+timedelta(days=1)
        a=self.change(kind='move',new_date=destination);svc.publish_change(a,self.admin)
        self.assertEqual(svc.schedule_on(self.day)[0]['status'],'moved_out')
        self.assertEqual(svc.schedule_on(destination)[0]['date'],destination.isoformat())
        b=self.change(kind='move',new_date=destination+timedelta(days=1));svc.publish_change(b,self.admin)
        self.assertEqual(svc.schedule_on(destination),[])
        self.assertEqual(len(svc.schedule_on(destination+timedelta(days=1))),1)
    def test_extra_lesson(self):
        change=self.change(kind='add',original=None,group=self.group,subject=self.subject,teacher=self.teacher,time_slot=self.slot2)
        svc.publish_change(change,self.admin)
        self.assertEqual(len(svc.schedule_on(self.day)),2)
    def test_room_teacher_group_conflicts(self):
        self.make_lesson()
        kinds={x['type'] for x in svc.conflicts_for(self.day,svc.schedule_on(self.day))}
        self.assertEqual(len(kinds),3)
    def test_distinct_subgroups_are_allowed(self):
        Participant.objects.filter(lesson=self.lesson).update(subgroup=self.sub1)
        self.make_lesson(subgroup=self.sub2,room=self.room2,teacher=self.teacher2)
        self.assertEqual(svc.conflicts_for(self.day,svc.schedule_on(self.day)),[])
    def test_whole_group_conflicts_with_subgroup(self):
        self.make_lesson(subgroup=self.sub2,room=self.room2,teacher=self.teacher2)
        self.assertTrue(any('группы' in x['type'] for x in svc.conflicts_for(self.day,svc.schedule_on(self.day))))
    def test_exact_time_overlap_not_slot_number(self):
        self.make_lesson(time_slot=self.slot2,start_time=time(9,45),end_time=time(10,30),room=self.room2,teacher=self.teacher2)
        self.assertTrue(svc.conflicts_for(self.day,svc.schedule_on(self.day)))
    def test_interval_boundaries_do_not_overlap(self):
        self.assertEqual(svc.availability(self.day,time(10),time(10,10),[self.room])[0]['status'],'free')
    def test_block_occupies_entire_interval(self):
        RoomBlock.objects.create(room=self.room2,start_at=timezone.make_aware(datetime.combine(self.day,time(9))),end_at=timezone.make_aware(datetime.combine(self.day,time(9,15))),reason='Ремонт',created_by=self.admin)
        self.assertEqual(svc.availability(self.day,time(8),time(10),[self.room2])[0]['status'],'unavailable')
    def test_unknown_is_not_free(self):
        self.assertEqual(svc.availability(date(2030,1,1),time(8),time(10),[self.room])[0]['status'],'unknown')
    def test_conflicting_change_cannot_publish(self):
        self.make_lesson(room=self.room2,teacher=self.teacher2,group=self.group2)
        change=self.change()
        with self.assertRaises(ValidationError):svc.publish_change(change,self.admin)
        change.refresh_from_db();self.assertEqual(change.status,'draft')
    def test_dispatcher_cannot_cross_scope(self):
        with self.assertRaises(PermissionDenied):svc.publish_change(self.change(room=self.remote),self.dispatcher)
    def test_dispatcher_can_publish_in_scope(self):
        svc.publish_change(self.change(),self.dispatcher)
        self.assertEqual(svc.schedule_on(self.day)[0]['room_id'],self.room2.pk)
    def test_holiday_and_override(self):
        calendar=CalendarDay.objects.create(year=self.year,date=self.day,is_holiday=True)
        self.assertEqual(svc.schedule_on(self.day),[])
        calendar.is_holiday=False;calendar.weekday_override=1;calendar.save()
        self.assertEqual(svc.schedule_on(self.day),[])
    def test_week_parity(self):
        self.lesson.week_type='odd';self.lesson.save()
        a=svc.schedule_on(self.day);b=svc.schedule_on(self.day+timedelta(days=7))
        self.assertNotEqual(bool(a),bool(b))
    def test_import_is_preview_and_confirm_atomic(self):
        batch=self.preview();self.assertEqual(batch.errors,[])
        self.assertEqual(ScheduleVersion.objects.count(),1)
        version=confirm_import(batch,self.admin)
        self.assertEqual(version.status,'published')
        self.assertEqual(svc.version_on(self.day).pk,version.pk)
        self.assertEqual(confirm_import(batch,self.admin).pk,version.pk)
    def test_import_unknown_values_do_not_write_schedule(self):
        file=self.workbook([['Понедельник',1,self.group.name,'','Unknown',self.teacher.name,self.building.name,self.room.number]])
        batch=self.preview(file);self.assertTrue(batch.errors)
        with self.assertRaises(ValidationError):confirm_import(batch,self.admin)
        self.assertEqual(ScheduleVersion.objects.count(),1)
    def test_duplicate_file_is_rejected(self):
        data=self.workbook().read()
        a=self.preview(SimpleUploadedFile('a.xlsx',data));confirm_import(a,self.admin)
        b=self.preview(SimpleUploadedFile('b.xlsx',data))
        with self.assertRaises(ValidationError):confirm_import(b,self.admin)
    def test_failed_import_rolls_back_all_rows(self):
        rows=[['Понедельник',1,self.group.name,'',self.subject.name,self.teacher.name,self.building.name,self.room.number],
              ['Понедельник',1,self.group2.name,'',self.subject.name,self.teacher2.name,self.building.name,self.room.number]]
        batch=self.preview(self.workbook(rows));self.assertEqual(batch.errors,[])
        with self.assertRaises(ValidationError):confirm_import(batch,self.admin)
        self.assertEqual(ScheduleVersion.objects.count(),1);self.assertEqual(Lesson.objects.count(),1)
    def test_merged_cells(self):
        book=Workbook();sheet=book.active;sheet.append([HEADERS[k] for k in FIELDS])
        sheet.append(['Понедельник',1,self.group.name,'',self.subject.name,self.teacher.name,self.building.name,self.room.number])
        sheet.append([None,2,self.group.name,'',self.subject.name,self.teacher.name,self.building.name,self.room.number])
        sheet.merge_cells('A2:A3');output=io.BytesIO();book.save(output)
        rows,errors=parse_xlsx(output.getvalue(),self.profile)
        self.assertEqual(rows[1]['day'],'Понедельник');self.assertEqual(errors,[])
    def test_bad_file(self):
        with self.assertRaises(ValidationError):self.preview(SimpleUploadedFile('fake.xlsx',b'not zip'))
        with self.assertRaises(ValidationError):self.preview(SimpleUploadedFile('file.exe',b'no'))
    def test_rollover_preserves_old_groups(self):
        target=AcademicYear.objects.create(name='2027/2028',start_date=date(2027,9,1),end_date=date(2028,8,31))
        preview=svc.rollover_preview(self.year,target)
        self.assertEqual(next(x for x in preview if x['group_id']==self.group.pk)['status'],'graduated')
        svc.rollover(self.year,target,self.admin)
        self.assertEqual(GroupYear.objects.filter(year=self.year).count(),2)
        svc.activate_year(target,self.admin)
        self.group.refresh_from_db();self.assertEqual(self.group.status,'graduated')
        self.assertEqual(len(svc.schedule_on(self.day)),1)
        with self.assertRaises(ValidationError):svc.rollover(self.year,target,self.admin)
    def test_api_search_cyrillic_and_initials(self):
        result=self.client.get('/api/teachers',{'q':'пиГАчева а. с.'}).json()
        self.assertEqual(result['count'],1)
    def test_public_pages_and_week(self):
        for url in ['/','/groups','/teachers','/rooms','/changes',f'/groups/{self.group.pk}']:
            self.assertEqual(self.client.get(url).status_code,200,url)
        result=self.client.get(f'/api/groups/{self.group.pk}/schedule',{'date':str(self.day),'view':'week'}).json()
        self.assertEqual(len(result['days']),6)
    def test_admin_api_requires_login_and_csrf(self):
        self.assertEqual(self.client.post('/api/admin/changes',data='{}',content_type='application/json').status_code,401)
        client=Client(enforce_csrf_checks=True);client.force_login(self.admin)
        self.assertEqual(client.post('/api/admin/changes',data='{}',content_type='application/json').status_code,403)
    def test_api_invalid_time_error(self):
        response=self.client.get('/api/rooms/availability',{'start':'15:00','end':'10:00'})
        self.assertEqual(response.status_code,400);self.assertIn('error',response.json())
    def test_admin_index_import_and_rollover(self):
        self.client.force_login(self.admin)
        for url in ['/admin/','/admin/schedule/importbatch/upload/','/admin/schedule/academicyear/rollover/','/admin/conflicts/']:
            self.assertEqual(self.client.get(url).status_code,200,url)
    def test_api_draft_and_publish(self):
        self.client.force_login(self.admin)
        response=self.client.post('/api/admin/changes',data=json.dumps({'target_date':str(self.day),'original_id':self.lesson.pk,'kind':'replace','room_id':self.room2.pk}),content_type='application/json')
        self.assertEqual(response.status_code,201,response.content)
        pk=response.json()['id']
        self.assertEqual(self.client.post(f'/api/admin/changes/{pk}/publish').status_code,200)
        response=self.client.patch(f'/api/admin/changes/{pk}',data='{}',content_type='application/json')
        self.assertEqual(response.status_code,400)
    def test_version_rollback(self):
        batch=self.preview();version=confirm_import(batch,self.admin)
        svc.rollback_version(version,self.admin)
        self.assertEqual(svc.version_on(self.day).pk,self.version.pk)
    def test_free_intervals(self):
        gaps=svc.free_intervals(svc.schedule_on(self.day),self.day)
        self.assertEqual(gaps,[{'start':'10:00','end':'11:40'}])
    def test_import_report_does_not_persist_preview(self):
        batch=self.preview()
        report=import_report(batch)
        self.assertEqual(report['conflicts'],[])
        self.assertEqual(report['rows'][0]['teacher'],self.teacher.name)
        self.assertEqual(ScheduleVersion.objects.count(),1)
        self.assertEqual(Lesson.objects.count(),1)
    def test_blocks_profile(self):
        profile=ImportProfile.objects.create(name='Блоки',kind='blocks',mapping={'group_row':1,'data_row':3,'day_column':1,'slot_column':2,
            'blocks':[{'group_column':3,'building':self.building.name,'columns':{'subject':3,'teacher':4,'room':5}}]})
        book=Workbook();sheet=book.active
        sheet.cell(1,3,self.group.name)
        sheet.append([])
        sheet.cell(3,1,'Понедельник');sheet.cell(3,2,1);sheet.cell(3,3,self.subject.name)
        sheet.cell(3,4,self.teacher.name);sheet.cell(3,5,self.room.number)
        output=io.BytesIO();book.save(output)
        batch=preview_import(SimpleUploadedFile('blocks.xlsx',output.getvalue()),profile,self.period,self.period.start_date,self.period.end_date,self.admin)
        self.assertEqual(batch.errors,[]);self.assertEqual(len(batch.payload),1)
    def test_teacher_reassignment_does_not_occupy_original_teacher(self):
        change=self.change(teacher=self.teacher2);svc.publish_change(change,self.admin)
        result=self.client.get(f'/api/teachers/{self.teacher.pk}/schedule',{'date':str(self.day)}).json()
        self.assertEqual(result['days'][0]['lessons'][0]['status'],'reassigned')
        self.assertEqual(result['days'][0]['free_intervals'],[{'start':'08:30','end':'11:40'}])
    def test_cancelled_import_cannot_confirm(self):
        batch=self.preview();batch.status='cancelled';batch.save()
        with self.assertRaises(ValidationError):confirm_import(batch,self.admin)
    def test_formula_cells_are_not_silently_evaluated(self):
        file=self.workbook([['Понедельник',1,self.group.name,'','=1+1',self.teacher.name,self.building.name,self.room.number]])
        with self.assertRaises(ValidationError):self.preview(file)
    def test_room_grid_has_all_slots(self):
        response=self.client.get('/api/rooms/availability',{'date':str(self.day),'start':'08:30','end':'10:00','grid':'1'})
        self.assertEqual(response.status_code,200,response.content)
        row=next(x for x in response.json()['results'] if x['id']==self.room.pk)
        self.assertEqual(row['grid'],[{'slot':1,'status':'busy'},{'slot':2,'status':'free'}])
    def test_dispatcher_cannot_import_over_another_building(self):
        self.make_lesson(room=self.remote,group=self.group2,teacher=self.teacher2,time_slot=self.slot2)
        batch=self.preview();batch.uploaded_by=self.dispatcher;batch.save()
        with self.assertRaises(PermissionDenied):confirm_import(batch,self.dispatcher)
        self.assertEqual(ScheduleVersion.objects.count(),1)
    def test_scope_enforced_on_api_patch(self):
        change=self.change()
        self.client.force_login(self.dispatcher)
        response=self.client.patch(f'/api/admin/changes/{change.pk}',data=json.dumps({'room_id':self.remote.pk}),content_type='application/json')
        self.assertEqual(response.status_code,403)
        change.refresh_from_db();self.assertEqual(change.room_id,self.room2.pk)
    def test_inactive_group_cannot_receive_extra_lesson(self):
        GroupYear.objects.filter(group=self.group).update(status='graduated')
        change=self.change(kind='add',original=None,group=self.group,subject=self.subject,teacher=self.teacher,time_slot=self.slot2)
        with self.assertRaises(ValidationError):svc.publish_change(change,self.admin)
    def test_published_lessons_cannot_be_edited_in_admin(self):
        self.client.force_login(self.admin)
        response=self.client.post(f'/admin/schedule/lesson/{self.lesson.pk}/change/',{})
        self.assertEqual(response.status_code,403)
    def test_backup_can_be_restored_to_separate_sqlite_file(self):
        import tempfile
        import sqlite3
        from contextlib import closing
        from pathlib import Path
        from django.core.management import call_command
        from django.test import override_settings
        with tempfile.TemporaryDirectory() as folder:
            source=Path(folder)/'source.sqlite3'
            with closing(sqlite3.connect(source)) as db:
                db.execute('CREATE TABLE sample (value TEXT)')
                db.execute("INSERT INTO sample VALUES ('retained')")
                db.commit()
            with override_settings(DATABASES={'default':{'ENGINE':'django.db.backends.sqlite3','NAME':source}}):
                call_command('backup_db',directory=str(Path(folder)/'copies'),keep=2,stdout=io.StringIO())
            backup=next((Path(folder)/'copies').glob('*.sqlite3'))
            with closing(sqlite3.connect(backup)) as db:
                self.assertEqual(db.execute('SELECT value FROM sample').fetchone()[0],'retained')
